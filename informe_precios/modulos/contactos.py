"""Gestión de la lista de contactos de envío (Excel "Contactos Envio Diario.xlsx").

El Excel es la fuente de verdad: hoja config.CONTACTOS_HOJA, columna B (índice
config.CONTACTOS_COL), cabecera en la fila 4 y destinatarios desde la fila
config.CONTACTOS_FILA_INICIO hacia abajo.
"""

import csv
import io
import logging
import re
from pathlib import Path

import openpyxl

import config

log = logging.getLogger(__name__)

_CORREO_RE = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")


def _normalizar_correo(valor):
    """Limpia un valor de celda a su forma de correo o None."""
    if valor is None:
        return None
    texto = str(valor).strip()
    if not texto:
        return None
    return texto.lower()


def leer_contactos():
    """Lee los correos de la hoja de contactos (columna B desde la fila 5)."""
    ruta = config.CONTACTOS_XLSX
    if not ruta.exists():
        log.warning("No existe %s, no hay destinatarios", ruta.name)
        return []

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[config.CONTACTOS_HOJA]

    contactos = []
    for fila in ws.iter_rows(min_row=config.CONTACTOS_FILA_INICIO,
                             max_col=config.CONTACTOS_COL):
        correo = _normalizar_correo(fila[config.CONTACTOS_COL - 1].value)
        if correo and correo not in contactos:
            contactos.append(correo)

    wb.close()
    log.info("Destinatarios (%d): %s", len(contactos), ", ".join(contactos))
    return contactos


def _ultima_fila_datos(ws):
    """Última fila con contenido en la columna de contactos (desde la inicial)."""
    ultima = config.CONTACTOS_FILA_INICIO - 1
    for fila in ws.iter_rows(min_row=config.CONTACTOS_FILA_INICIO,
                             max_col=config.CONTACTOS_COL):
        valor = _normalizar_correo(fila[config.CONTACTOS_COL - 1].value)
        if valor:
            ultima = fila[0].row
    return ultima


def guardar_contactos(emails):
    """Sobrescribe la lista de destinatarios en el Excel.

    Conserva la cabecera (fila 4) y el resto del libro. Limpia los sobrantes
    de guardados anteriores. Devuelve la lista normalizada final.
    """
    emails = _depurar_emails(emails)
    ruta = config.CONTACTOS_XLSX

    wb = openpyxl.load_workbook(ruta)
    ws = wb[config.CONTACTOS_HOJA]
    col = config.CONTACTOS_COL

    # Limpiar celdas sobrantes (guardados previos con más filas)
    ultima = _ultima_fila_datos(ws)
    for fila in range(config.CONTACTOS_FILA_INICIO + len(emails), ultima + 1):
        ws.cell(row=fila, column=col).value = None

    for i, correo in enumerate(emails):
        ws.cell(row=config.CONTACTOS_FILA_INICIO + i, column=col).value = correo

    wb.save(ruta)
    log.info("Guardados %d destinatarios en %s", len(emails), ruta.name)
    return emails


def _depurar_emails(emails):
    """Normaliza, elimina vacíos y duplicados de una lista de correos."""
    vistos = []
    for v in emails:
        correo = _normalizar_correo(v)
        if correo and correo not in vistos:
            vistos.append(correo)
    return vistos


def anadir_contactos(emails):
    """Añade correos a la lista existente sin duplicados y persiste.

    Devuelve la lista final completa.
    """
    emails = _depurar_emails(emails)
    if not emails:
        return leer_contactos()
    finales = _depurar_emails(leer_contactos() + emails)
    return guardar_contactos(finales)


def eliminar_contactos(emails):
    """Elimina los correos indicados de la lista y persiste.

    Devuelve la lista final completa.
    """
    a_borrar = {_normalizar_correo(e) for e in emails if _normalizar_correo(e)}
    if not a_borrar:
        return leer_contactos()
    finales = [c for c in leer_contactos() if c not in a_borrar]
    return guardar_contactos(finales)


def _celdas_de_xlsx(ruta):
    """Devuelve los textos no vacíos de todas las celdas de un .xlsx."""
    celdas = []
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for fila in ws.iter_rows():
                for celda in fila:
                    valor = celda.value
                    if valor is not None and str(valor).strip():
                        celdas.append(str(valor))
    finally:
        wb.close()
    return celdas


def _celdas_de_csv(ruta):
    """Devuelve los textos no vacíos de todas las celdas de un .csv."""
    celdas = []
    datos = ruta.read_bytes()
    for codif in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            texto = datos.decode(codif)
            break
        except UnicodeDecodeError:
            continue
    for fila in csv.reader(io.StringIO(texto)):
        for celda in fila:
            if celda.strip():
                celdas.append(celda)
    return celdas


def leer_emails_archivo(ruta):
    """Extrae los correos válidos de un .xlsx o .csv (cualquier layout).

    El archivo puede tener la lista en la columna que sea, con o sin cabecera:
    se escanean todas las celdas buscando patrones de correo. Devuelve la
    lista normalizada sin duplicados.
    """
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".csv":
        celdas = _celdas_de_csv(ruta)
    else:
        celdas = _celdas_de_xlsx(ruta)

    vistos = []
    for celda in celdas:
        for match in _CORREO_RE.findall(celda):
            correo = _normalizar_correo(match)
            if correo and correo not in vistos:
                vistos.append(correo)
    log.info("Archivo %s: %d celdas con datos, %d correos extraídos",
             ruta.name, len(celdas), len(vistos))
    return vistos


def exportar_excel_bytes():
    """Devuelve un .xlsx en memoria con una sola columna de correos."""
    stream = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Correos"
    ws.cell(row=1, column=1, value="Correo")
    ws.column_dimensions["A"].width = 40
    for i, correo in enumerate(leer_contactos(), start=2):
        ws.cell(row=i, column=1, value=correo)
    wb.save(stream)
    stream.seek(0)
    return stream
